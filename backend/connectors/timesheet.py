"""Timesheet (beta): fills the monthly sheets of the Excel timesheet from a plain description.

The workbook (TIMESHEET_PATH) has one sheet per month, "MM_YYYY", with a row per day
from row 12: C/D morning start/end, E/F afternoon start/end, H the type of day (a
dropdown fed by Listen!L1:L14), I the location. Everything else (weekday names, hours,
totals, the yearly sheet and its charts) is calculated by Excel and never touched.

Saving never goes through an Excel library: those drop the dropdowns and can damage
the charts. Only the XML of the changed cells is rewritten inside the .xlsx; every
other part of the file stays byte-for-byte the same. A backup is written next to the
file first, and Excel is told to recalculate the totals when the file is opened.

Public holidays are the day rows the template colours pink-red; they are filled as
"Feriado - Feiertag". Rows coloured green are asked about rather than guessed.
"""
from __future__ import annotations

import calendar
import datetime as dt
import html
import os
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

FIRST_ROW = 12
HOLIDAY_FILL = "FFFADCEB"   # pink-red: public holidays in the template
SPECIAL_FILL = "FFDAEDE6"   # green: e.g. 24 and 31 December
WEEKDAYS_PT = ["Seg.", "Ter.", "Qua.", "Qui.", "Sex.", "Sáb.", "Dom."]

WORK = "Trabalho - Istarbeit"
HOLIDAY = "Feriado - Feiertag"
# Day types that have working hours and a location; every other type (vacation, sick
# leave, absences...) has neither, as in the months filled so far.
WITH_HOURS = {WORK, "Trabalho - Istarbeit (Nichtauslastung)", "Treinamento - Training", "KTP - KTP"}
# Plain-English names for the sheet's day types, for what J.A.R.V.I.G. says and shows.
FRIENDLY = {
    "trabalho - istarbeit": "Work",
    "trabalho - istarbeit (nichtauslastung)": "Work (not utilised)",
    "treinamento - training": "Training",
    "férias - urlaub": "Vacation",
    "feriado - feiertag": "Public holiday",
    "baixa médica - krankschreibung arzt": "Sick leave",
    "baixa seguro - krankschreibung arbeitsunfall": "Sick leave (work accident)",
    "baixa gravidez de risco - krankschreibung wegen risikoschwangerschaft": "Sick leave (risk pregnancy)",
    "falta justificada - gerechtfertigte fehlzeit": "Justified absence",
    "falta injustificada - ungerechtfertigte fehlzeit": "Unjustified absence",
    "assistência à familia - unterstützung für familien": "Family assistance",
    "licença de parentalidade - elternzeit": "Parental leave",
    "ktp - ktp": "KTP",
}


def friendly(day_type: str | None) -> str:
    return FRIENDLY.get((day_type or "").strip().lower(), (day_type or "").strip())


MONTHS_PT = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
             "setembro", "outubro", "novembro", "dezembro"]


class TimesheetError(Exception):
    """A problem the user has to fix (missing file, missing month, file open in Excel...)."""


def _hhmm(value) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (dt.time, dt.datetime)):
        return value.strftime("%H:%M")
    if isinstance(value, (int, float)):  # Excel fraction of a day
        minutes = round(value * 24 * 60)
        return f"{minutes // 60:02d}:{minutes % 60:02d}"
    return str(value)


def parse_time(text: str) -> dt.time:
    m = re.fullmatch(r"\s*(\d{1,2})(?:[:h.](\d{2}))?\s*", str(text))
    if not m or int(m.group(1)) > 23 or int(m.group(2) or 0) > 59:
        raise ValueError(f"'{text}' is not a time")
    return dt.time(int(m.group(1)), int(m.group(2) or 0))


# ------------------------------------------------------------------ reading
@dataclass
class DayRow:
    day: int
    row: int
    weekday: int             # 0 = Monday
    holiday: bool            # pink-red row in the template
    special: bool            # green row in the template
    current: dict            # {"C".."F": "HH:MM"|None, "H": str|None, "I": str|None}

    @property
    def weekend(self) -> bool:
        return self.weekday >= 5

    @property
    def filled(self) -> bool:
        return any(self.current.values())


@dataclass
class Month:
    year: int
    month: int
    sheet: str
    days: list[DayRow]
    day_types: list[str]     # the H dropdown values
    locations: list[str]
    holiday_source: tuple[str, int] | None = None  # (sheet, row) already coloured as a holiday, to copy

    @property
    def label(self) -> str:
        return f"{calendar.month_name[self.month]} {self.year}"


def read_month(path: Path, year: int, month: int) -> Month:
    """What the month's sheet contains now, plus the allowed values. Read-only."""
    if not path.exists():
        raise TimesheetError(f"I can't find the timesheet at {path}. Check TIMESHEET_PATH in backend/.env.")
    wb = openpyxl.load_workbook(path)  # reading only: this workbook object is never saved
    sheet = f"{month:02d}_{year}"
    if sheet not in wb.sheetnames:
        raise TimesheetError(f"The timesheet has no sheet for {calendar.month_name[month]} {year} ('{sheet}').")
    ws = wb[sheet]
    listen = wb["Listen"] if "Listen" in wb.sheetnames else None
    day_types = [str(c.value) for c in listen["L"][:20] if c.value] if listen else [WORK, HOLIDAY]
    locations = [str(c.value) for c in listen["P"][:10] if c.value] if listen else ["Portugal"]
    days = []
    for day in range(1, calendar.monthrange(year, month)[1] + 1):
        row = FIRST_ROW + day - 1
        fill = ws[f"A{row}"].fill
        rgb = fill.fgColor.rgb if fill.fill_type == "solid" and isinstance(fill.fgColor.rgb, str) else None
        current = {c: _hhmm(ws[f"{c}{row}"].value) for c in "CDEF"}
        current |= {c: (str(ws[f"{c}{row}"].value) if ws[f"{c}{row}"].value else None) for c in "HI"}
        days.append(DayRow(day, row, dt.date(year, month, day).weekday(), rgb == HOLIDAY_FILL,
                           rgb == SPECIAL_FILL, current))
    source = next(((sheet, d.row) for d in days if d.holiday), None)
    for name in (n for n in wb.sheetnames if re.fullmatch(r"\d{2}_\d{4}", n) and source is None):
        other = wb[name]
        for row in range(FIRST_ROW, FIRST_ROW + 31):
            fill = other[f"A{row}"].fill
            if fill.fill_type == "solid" and fill.fgColor.rgb == HOLIDAY_FILL:
                source = (name, row)
                break
    return Month(year, month, sheet, days, day_types, locations, source)


# ------------------------------------------------------------------ planning
@dataclass
class DayPlan:
    day: int
    type: str
    morning: tuple[str, str] | None = None
    afternoon: tuple[str, str] | None = None
    location: str | None = None
    reason: str = ""          # "work", "holiday", "said"

    def cells(self) -> dict:
        return {"C": self.morning[0] if self.morning else None, "D": self.morning[1] if self.morning else None,
                "E": self.afternoon[0] if self.afternoon else None, "F": self.afternoon[1] if self.afternoon else None,
                "H": self.type, "I": self.location}


def split_hours(start: str, end: str, lunch: tuple[str, str]) -> tuple[tuple[str, str] | None, tuple[str, str] | None]:
    """8:00-17:00 with lunch 12:00-13:00 -> (08:00, 12:00), (13:00, 17:00). Half days keep one part."""
    s, e = parse_time(start), parse_time(end)
    ls, le = parse_time(lunch[0]), parse_time(lunch[1])
    if e <= s:
        raise ValueError(f"the end time {end} is not after the start time {start}")
    fmt = lambda t: t.strftime("%H:%M")
    morning = (fmt(s), fmt(min(e, ls))) if s < ls else None
    afternoon = (fmt(max(s, le)), fmt(e)) if e > le else None
    if not morning and not afternoon:  # entirely inside the lunch break
        morning = (fmt(s), fmt(e))
    return morning, afternoon


@dataclass
class Plan:
    month: Month
    days: dict[int, DayPlan] = field(default_factory=dict)
    questions: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def changes(self) -> dict[int, DayPlan]:
        """Days whose cells would actually change."""
        rows = {d.day: d for d in self.month.days}
        return {n: p for n, p in self.days.items() if p.cells() != rows[n].current}

    def overwritten(self) -> list[int]:
        rows = {d.day: d for d in self.month.days}
        return [n for n in self.changes() if rows[n].filled]


def build_plan(month: Month, spec: dict, lunch: tuple[str, str], default_location: str) -> Plan:
    """Turn the structured description (from the AI) into exact cell values, with the sheet's rules:
    weekends stay empty unless named, pink-red rows become holidays, green rows are asked about."""
    plan = Plan(month)
    types = {t.strip().lower(): t for t in month.day_types}
    locations = {loc.lower(): loc for loc in month.locations}
    by_day = {d.day: d for d in month.days}

    def canonical_type(value: str | None) -> str | None:
        if not value:
            return None
        found = types.get(str(value).strip().lower())
        if not found:
            plan.questions.append(f"I don't know the day type '{value}'. The sheet allows: {', '.join(month.day_types)}.")
        return found

    def canonical_location(value: str | None) -> str | None:
        if not value:
            return None
        aliases = {"germany": "alemanha", "deutschland": "alemanha"}
        found = locations.get(aliases.get(str(value).lower(), str(value).lower()))
        if not found:
            plan.questions.append(f"'{value}' isn't one of the sheet's locations ({', '.join(month.locations)}).")
        return found

    def make(day: int, type_: str, start: str | None, end: str | None, location: str | None, reason: str):
        if type_ in WITH_HOURS:
            if not (start and end):
                plan.questions.append(f"What hours did you work on the {day}{_suffix(day)}?")
                return
            try:
                morning, afternoon = split_hours(start, end, lunch)
            except ValueError as exc:
                plan.questions.append(f"For the {day}{_suffix(day)}: {exc}.")
                return
            plan.days[day] = DayPlan(day, type_, morning, afternoon, location or default_location, reason)
        elif type_ == HOLIDAY:
            plan.days[day] = DayPlan(day, type_, location=location or default_location, reason=reason)
        else:  # vacation, sick leave, absences: no hours, no location
            plan.days[day] = DayPlan(day, type_, reason=reason)

    work = spec.get("work") or None
    work_type = canonical_type((work or {}).get("type")) or WORK
    work_location = canonical_location((work or {}).get("location"))
    overrides = {}
    for item in spec.get("days") or []:
        try:
            n = int(item["day"])
        except (KeyError, TypeError, ValueError):
            continue
        if n not in by_day:
            plan.warnings.append(f"{calendar.month_name[month.month]} has no day {n}; I left it out.")
            continue
        overrides[n] = item

    for d in month.days:
        item = overrides.get(d.day)
        if item is not None:
            type_ = canonical_type(item.get("type")) or work_type
            start = item.get("start") or ((work or {}).get("start") if type_ in WITH_HOURS else None)
            end = item.get("end") or ((work or {}).get("end") if type_ in WITH_HOURS else None)
            make(d.day, type_, start, end, canonical_location(item.get("location")) or work_location, "said")
            if d.weekend:
                plan.warnings.append(f"The {d.day}{_suffix(d.day)} is a {calendar.day_name[d.weekday]}; "
                                     "I've filled it as you said.")
        elif d.holiday:
            make(d.day, HOLIDAY, None, None, work_location, "holiday")
        elif d.weekend or not work:
            continue
        elif d.special:
            plan.questions.append(f"The {d.day}{_suffix(d.day)} is marked green in the template. "
                                  "Was it a normal working day, or something else?")
        else:
            make(d.day, work_type, work.get("start"), work.get("end"), work_location, "work")

    for q in spec.get("questions") or []:
        if q and q not in plan.questions:
            plan.questions.append(str(q))
    return plan


def _suffix(n: int) -> str:
    return "th" if 11 <= n % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")


def _ranges(days: list[int]) -> str:
    days = sorted(days)
    out, start = [], None
    for i, d in enumerate(days):
        if start is None:
            start = d
        if i + 1 == len(days) or days[i + 1] != d + 1:
            out.append(str(start) if start == d else f"{start}–{d}")
            start = None
    return ", ".join(out)


def describe(plan: Plan) -> tuple[str, str]:
    """(short text summary for the conversation, shorter version to read aloud)."""
    groups: dict[tuple, list[int]] = {}
    for n, p in sorted(plan.days.items()):
        hours = " + ".join(f"{part[0]}–{part[1]}" for part in (p.morning, p.afternoon) if part)
        groups.setdefault((friendly(p.type), hours, p.location or ""), []).append(n)
    lines, spoken = [], []
    for (name, hours, location), days in groups.items():
        detail = ", ".join(x for x in (hours, location) if x)
        lines.append(f"• {name}{' (' + detail + ')' if detail else ''}: {_ranges(days)}")
        spoken.append(f"{len(days)} day{'s' if len(days) != 1 else ''} of {name.lower()}")
    weekends = [d.day for d in plan.month.days if d.weekend and d.day not in plan.days]
    if weekends:
        lines.append(f"• Weekends left empty: {_ranges(weekends)}")
    overwritten = plan.overwritten()
    if overwritten:
        lines.append(f"• Already filled, will be replaced: {_ranges(overwritten)}")
    lines += [f"• Note: {w}" for w in plan.warnings]
    say = (f"For {plan.month.label}: " + ", ".join(spoken) + ".") if spoken else f"Nothing to change in {plan.month.label}."
    if overwritten:
        say += f" {len(overwritten)} day{'s' if len(overwritten) != 1 else ''} already filled will be replaced."
    return "\n".join(lines), say


def table(plan: Plan) -> dict:
    """Day-by-day table of exactly what goes into the sheet, for the HUD."""
    rows = []
    changes = plan.changes()
    for d in plan.month.days:
        p = plan.days.get(d.day)
        cells = p.cells() if p else d.current
        hours = 0
        for a, b in ((cells["C"], cells["D"]), (cells["E"], cells["F"])):
            if a and b:
                ta, tb = parse_time(a), parse_time(b)
                hours += (tb.hour * 60 + tb.minute - ta.hour * 60 - ta.minute) / 60
        if d.day in changes:
            status = "replace" if d.filled else "new"
        else:
            status = "same" if d.filled else ""
        kind = "holiday" if (cells["H"] or "").strip().lower() == HOLIDAY.lower() or d.holiday else \
            "weekend" if d.weekend else "special" if d.special else ""
        short = lambda t: t.lstrip("0") if t and not t.startswith("00") else t  # "08:00" -> "8:00"
        rows.append({
            "cells": [str(d.day), WEEKDAYS_PT[d.weekday],
                      f"{short(cells['C'])}–{short(cells['D'])}" if cells["C"] else "",
                      f"{short(cells['E'])}–{short(cells['F'])}" if cells["E"] else "",
                      (f"{hours:g} h" if hours else ""), friendly(cells["H"]), cells["I"] or "",
                      {"new": "new", "replace": "replaces", "same": "unchanged", "": ""}[status]],
            "kind": kind, "status": status,
        })
    return {"title": f"{plan.month.label} · sheet {plan.month.sheet}",
            "columns": ["Day", "", "Morning", "Afternoon", "Hours", "Type", "Location", ""],
            "rows": rows}


# ------------------------------------------------------------------ saving (cell XML only)
def _excel_time(hhmm: str) -> str:
    t = parse_time(hhmm)
    return repr((t.hour * 60 + t.minute) / (24 * 60))


def _sheet_part(z: zipfile.ZipFile, sheet: str) -> str:
    wbx = z.read("xl/workbook.xml").decode()
    m = re.search(rf'<sheet [^>]*name="{re.escape(html.escape(sheet))}"[^>]*r:id="(rId\d+)"', wbx)
    if not m:
        raise TimesheetError(f"Sheet '{sheet}' is missing from the workbook.")
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    target = re.search(rf'<Relationship [^>]*Id="{m.group(1)}"[^>]*Target="([^"]+)"', rels) or \
        re.search(rf'<Relationship [^>]*Target="([^"]+)"[^>]*Id="{m.group(1)}"', rels)
    path = target.group(1).lstrip("/")
    return path if path.startswith("xl/") else "xl/" + path


def _shared_strings(z: zipfile.ZipFile) -> dict[str, int]:
    try:
        xml = z.read("xl/sharedStrings.xml").decode()
    except KeyError:
        return {}
    index = {}
    for i, si in enumerate(re.findall(r"<si>(.*?)</si>", xml, re.S)):
        texts = re.findall(r"<t(?: [^>]*)?>(.*?)</t>", si, re.S)
        index.setdefault(html.unescape("".join(texts)), i)
    return index


_CELL = r'<c r="{ref}"(?P<attrs>(?: [^>]*?)?)(?:/>|>(?P<body>.*?)</c>)'


def _set_cell(row_xml: str, ref: str, value, strings: dict[str, int], style: str | None = None) -> str:
    """Rewrite one cell inside a <row>, keeping its style unless `style` is given."""
    m = re.search(_CELL.format(ref=ref), row_xml, re.S)
    attrs = m.group("attrs") if m else ""
    s = style or (re.search(r' s="(\d+)"', attrs).group(1) if re.search(r' s="(\d+)"', attrs) else None)
    s_attr = f' s="{s}"' if s else ""
    if value is None:
        cell = f'<c r="{ref}"{s_attr}/>'
    elif isinstance(value, str) and not re.fullmatch(r"[0-9.eE-]+", value):
        if value in strings:
            cell = f'<c r="{ref}"{s_attr} t="s"><v>{strings[value]}</v></c>'
        else:
            cell = f'<c r="{ref}"{s_attr} t="inlineStr"><is><t xml:space="preserve">{html.escape(value)}</t></is></c>'
    else:
        cell = f'<c r="{ref}"{s_attr}><v>{value}</v></c>'
    if m:
        return row_xml[:m.start()] + cell + row_xml[m.end():]
    # The cell doesn't exist yet: insert it in column order.
    col = re.match(r"[A-Z]+", ref).group(0)
    for later in re.finditer(r'<c r="([A-Z]+)\d+"', row_xml):
        if (len(later.group(1)), later.group(1)) > (len(col), col):
            return row_xml[:later.start()] + cell + row_xml[later.start():]
    return row_xml.replace("</row>", cell + "</row>")


def _restyle(row_xml: str, ref: str, style: str) -> str:
    m = re.search(rf'<c r="{ref}"(?P<attrs>(?: [^>]*?)?)(/?>)', row_xml)
    if not m:
        return row_xml
    attrs = re.sub(r' s="\d+"', "", m.group("attrs")) + f' s="{style}"'
    return row_xml[:m.start()] + f'<c r="{ref}"{attrs}{m.group(2)}' + row_xml[m.end():]


def _row_styles(sheet_xml: str, row: int) -> dict[str, str]:
    """{column: style id} for cells A..I of one row."""
    m = re.search(rf'<row r="{row}"[^>]*>(.*?)</row>', sheet_xml, re.S)
    return dict(re.findall(r'<c r="([A-I])\d+"[^>]*? s="(\d+)"', m.group(1))) if m else {}


def excel_lock(path: Path) -> Path | None:
    """Excel's '~$' lock file next to the workbook, if it is open."""
    for f in path.parent.glob("~$*"):
        if path.name[2:] in f.name or f.name[2:] in path.name:
            return f
    return None


def save(path: Path, plan: Plan) -> Path:
    """Write the plan's changed days into the workbook. Returns the backup path."""
    if excel_lock(path):
        raise TimesheetError("The timesheet seems to be open in Excel. Close it first, then say yes again.")
    changes = plan.changes()
    if not changes:
        raise TimesheetError("There's nothing to change.")
    with zipfile.ZipFile(path) as z:
        part = _sheet_part(z, plan.month.sheet)
        sheet_xml = z.read(part).decode()
        strings = _shared_strings(z)
        red = None
        if plan.month.holiday_source:
            src_sheet, src_row = plan.month.holiday_source
            src_xml = sheet_xml if src_sheet == plan.month.sheet else z.read(_sheet_part(z, src_sheet)).decode()
            red = _row_styles(src_xml, src_row) or None
        by_day = {d.day: d for d in plan.month.days}
        for n, day_plan in changes.items():
            row = by_day[n].row
            m = re.search(rf'<row r="{row}"[^>]*>.*?</row>', sheet_xml, re.S)
            if not m:
                raise TimesheetError(f"Row {row} for day {n} is missing from sheet {plan.month.sheet}.")
            row_xml = m.group(0)
            cells = day_plan.cells()
            for col in "CDEF":
                row_xml = _set_cell(row_xml, f"{col}{row}", _excel_time(cells[col]) if cells[col] else None, strings)
            for col in "HI":
                row_xml = _set_cell(row_xml, f"{col}{row}", cells[col], strings)
            if day_plan.type == HOLIDAY and not by_day[n].holiday and red:
                for col, style in red.items():  # colour it like the template's other holidays
                    row_xml = _restyle(row_xml, f"{col}{row}", style)
            sheet_xml = sheet_xml[:m.start()] + row_xml + sheet_xml[m.end():]
        workbook_xml = z.read("xl/workbook.xml").decode()
        # Cached totals are now stale: have Excel recalculate everything when the file is opened.
        if "fullCalcOnLoad" not in workbook_xml:
            workbook_xml = re.sub(r"<calcPr\b", '<calcPr fullCalcOnLoad="1"', workbook_xml, count=1) \
                if "<calcPr" in workbook_xml else workbook_xml.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')
        backup = path.with_name(f"{path.stem}.backup-{dt.datetime.now():%Y%m%d-%H%M%S}{path.suffix}")
        shutil.copy2(path, backup)
        fd, tmp = tempfile.mkstemp(suffix=path.suffix, dir=path.parent)
        os.close(fd)
        try:
            with zipfile.ZipFile(tmp, "w") as out:
                for info in z.infolist():  # same order, same compression; only two parts change
                    data = z.read(info.filename)
                    if info.filename == part:
                        data = sheet_xml.encode()
                    elif info.filename == "xl/workbook.xml":
                        data = workbook_xml.encode()
                    out.writestr(info, data, compress_type=info.compress_type)
        except Exception:
            os.unlink(tmp)
            raise
    os.replace(tmp, path)
    return backup
